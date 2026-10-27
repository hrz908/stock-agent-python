"""文档加载器：按文件类型分流解析为 Markdown 文本。"""
import hashlib
import logging
from pathlib import Path

logger = logging.getLogger(__name__)

MIME_MAP = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".md": "text/markdown",
    ".txt": "text/plain",
    ".csv": "text/csv",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def content_hash(content: str) -> str:
    """SHA-256 哈希用于幂等判断"""
    return hashlib.sha256(content.encode()).hexdigest()


# ------------------------------------------------------------------
# PDF
# ------------------------------------------------------------------

def _load_pdf(path: Path) -> str:
    """PDF 解析：优先 pymupdf4llm，失败降级 docling"""
    # 1) pymupdf4llm
    try:
        import pymupdf4llm
        md = pymupdf4llm.to_markdown(str(path))
        if md and md.strip():
            return md
        logger.warning("pymupdf4llm 返回空内容（可能是扫描件），尝试 docling: %s", path)
    except ImportError:
        logger.debug("pymupdf4llm 未安装，尝试 docling: %s", path)
    except Exception as e:
        logger.warning("pymupdf4llm 解析失败 (%s)，尝试 docling: %s", e, path)

    # 2) docling 降级
    try:
        from docling.document_converter import DocumentConverter
        converter = DocumentConverter()
        result = converter.convert(str(path))
        return result.document.export_to_markdown()
    except ImportError:
        logger.debug("docling 未安装，尝试 pymupdf 纯文本提取: %s", path)
    except Exception as e:
        logger.warning("docling 解析失败 (%s)，尝试纯文本提取: %s", e, path)

    # 3) 最终降级：pymupdf 纯文本
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(str(path))
        parts = []
        for page in doc:
            parts.append(page.get_text())
        doc.close()
        text = "\n\n".join(parts)
        if text.strip():
            return text
    except Exception as e:
        logger.error("所有 PDF 解析方式均失败: %s - %s", path, e)

    raise RuntimeError(f"无法解析 PDF 文件: {path}")


# ------------------------------------------------------------------
# DOCX
# ------------------------------------------------------------------

def _load_docx(path: Path) -> str:
    """DOCX 解析：优先 docling，降级 python-docx"""
    try:
        from docling.document_converter import DocumentConverter
        converter = DocumentConverter()
        result = converter.convert(str(path))
        return result.document.export_to_markdown()
    except ImportError:
        logger.debug("docling 未安装，使用 python-docx: %s", path)
    except Exception as e:
        logger.warning("docling 解析 DOCX 失败 (%s)，使用 python-docx: %s", e, path)

    try:
        import docx
        doc = docx.Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs)
    except Exception as e:
        raise RuntimeError(f"无法解析 DOCX 文件: {path} - {e}") from e


# ------------------------------------------------------------------
# PPTX
# ------------------------------------------------------------------

def _load_pptx(path: Path) -> str:
    """PPTX 解析：优先 docling，降级 python-pptx"""
    try:
        from docling.document_converter import DocumentConverter
        converter = DocumentConverter()
        result = converter.convert(str(path))
        return result.document.export_to_markdown()
    except ImportError:
        logger.debug("docling 未安装，使用 python-pptx: %s", path)
    except Exception as e:
        logger.warning("docling 解析 PPTX 失败 (%s)，使用 python-pptx: %s", e, path)

    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        slides_text = []
        for i, slide in enumerate(prs.slides, 1):
            texts = [shape.text for shape in slide.shapes if shape.has_text_frame]
            slide_md = f"## Slide {i}\n\n" + "\n\n".join(texts)
            slides_text.append(slide_md)
        return "\n\n---\n\n".join(slides_text)
    except ImportError as exc:
        raise RuntimeError(f"python-pptx 未安装，无法解析 PPTX: {path}") from exc
    except Exception as e:
        raise RuntimeError(f"无法解析 PPTX 文件: {path} - {e}") from e


# ------------------------------------------------------------------
# MD / TXT
# ------------------------------------------------------------------

def _load_text(path: Path) -> str:
    """直接读取纯文本/Markdown 文件"""
    encodings = ["utf-8", "gbk", "gb2312", "latin-1"]
    for enc in encodings:
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, LookupError):
            continue
    raise RuntimeError(f"无法解码文本文件: {path}")


# ------------------------------------------------------------------
# CSV / XLSX
# ------------------------------------------------------------------

def _load_csv(path: Path) -> str:
    """CSV 转 Markdown 表格"""
    import pandas as pd
    df = pd.read_csv(str(path))
    return df.to_markdown(index=False)


def _load_xlsx(path: Path) -> str:
    """XLSX 转 Markdown 表格（多 sheet 拼接）"""
    import pandas as pd
    xls = pd.ExcelFile(str(path))
    parts = []
    for sheet in xls.sheet_names:
        df = pd.read_excel(xls, sheet_name=sheet)
        parts.append(f"## {sheet}\n\n{df.to_markdown(index=False)}")
    return "\n\n".join(parts)


# ------------------------------------------------------------------
# 主入口
# ------------------------------------------------------------------

def load_document(path: str) -> tuple[str, str]:
    """
    加载文档并返回 (content_markdown, mime_type)。
    
    支持的格式: .pdf, .docx, .pptx, .md, .txt, .csv, .xlsx
    """
    p = Path(path)
    ext = p.suffix.lower()

    if ext == ".pdf":
        content = _load_pdf(p)
    elif ext == ".docx":
        content = _load_docx(p)
    elif ext == ".pptx":
        content = _load_pptx(p)
    elif ext in (".md", ".txt"):
        content = _load_text(p)
    elif ext == ".csv":
        content = _load_csv(p)
    elif ext == ".xlsx":
        content = _load_xlsx(p)
    else:
        raise ValueError(f"不支持的文件类型: {ext}")

    mime_type = MIME_MAP.get(ext, "application/octet-stream")
    return content, mime_type
